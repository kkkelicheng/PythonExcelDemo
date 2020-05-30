import openpyxl
from openpyxl.chart import(
	Reference,
	Series,
	PieChart,
	BarChart,
	BubbleChart
)
wb = openpyxl.Workbook()
ws_chat_pie = wb.active
ws_chat_pie.title = 'PieChart'
pie_data = [
    ['Pie', 'Sold'],
    ['Apple', 50],
    ['Cherry', 30],
    ['Pumpkin', 10],
    ['Chocolate', 40],
]
for row in pie_data:
    ws_chat_pie.append(row)

# 生成一个图表对象
pie = PieChart()
# 将Apple，Cherry ... Chocolate作为【指标】引用一下
labels = Reference(worksheet=ws_chat_pie,min_col=1,min_row=2,max_row=5)
# 将Apple，Cherry ... Chocolate对应的数据引用一下,注意一下，下面的是row 从1 开始
data = Reference(worksheet=ws_chat_pie,min_col=2,min_row=1,max_row=5)
pie.add_data(data,titles_from_data=True)
pie.set_categories(labels)
pie.title = 'Pies sold by category'
# 吧生成的图表加入sheet
ws_chat_pie.add_chart(pie,'A13')


# Bar chart
ws_barChart = wb.create_sheet('barChart')
rows = [
    ('Number', 'Batch 1', 'Batch 2'),
    (2, 10, 30),
    (3, 40, 60),
    (4, 50, 70),
    (5, 20, 10),
    (6, 10, 40),
    (7, 50, 30),
]

for row in rows:
    ws_barChart.append(row)

chart_bar = BarChart()
chart_bar.type = 'col' #竖直的， 不是横着的
chart_bar.style = 10 #
chart_bar.title = 'Bar Chart'
chart_bar.y_axis.title = 'Sample length(mm)'
chart_bar.x_axis.title = 'Test number(h)'
# bar chart 的 category 就是横坐标的那些
cate = Reference(ws_barChart, min_col=1, min_row=2, max_row=7)

# 注意一下，下面的y data的row是从1开始的，吧 名字 都算进去了，因为 cate 是row 2:7 所以会智能的去取名字
y_data = Reference(ws_barChart, min_col=2, max_col=3, min_row=1, max_row=7)
chart_bar.add_data(y_data,titles_from_data=True)
chart_bar.set_categories(cate)
chart_bar.shape = 4
# 吧生成的图表加入sheet
ws_barChart.add_chart(chart_bar,'A13')


wb.save(filename='chartExcel.xlsx')

